import time

from omr_studio.models import RunOptions, SessionState, SheetStatus
from omr_studio.paths import collect_sheets
from omr_studio.services.engine import EngineHooks, OMREngine, RunCancelled
from omr_studio.services.exporter import ResultsExporter
from omr_studio.services.results import ResultsReader
from omr_studio.services.sessions import SessionStore

import pytest
from openpyxl import load_workbook


def _run(workspace, library, samples, sheets, name="Keyed", hooks=None):
    template = library.import_template(samples / "answer-key" / "using-csv", name)
    store = SessionStore(workspace.sessions)
    options = RunOptions()
    session = store.create(template, sheets, options)
    started = time.monotonic()
    OMREngine().execute(session, options, hooks or EngineHooks())
    session = store.finish(session, SessionState.COMPLETED, time.monotonic() - started)
    return ResultsReader().read(session)


def test_graded_run_and_bad_file_isolation(workspace, library, samples, tmp_path):
    corrupt = tmp_path / "corrupt.jpg"
    corrupt.write_bytes(b"definitely not an image")
    sheets = collect_sheets([samples / "answer-key" / "using-csv" / "adrian_omr.png", corrupt])
    result = _run(workspace, library, samples, sheets)

    by_id = {s.file_id: s for s in result.sheets}
    assert by_id["adrian_omr.png"].status is SheetStatus.PROCESSED and by_id["adrian_omr.png"].score == 5.0
    assert by_id["adrian_omr.png"].marked_path and by_id["adrian_omr.png"].marked_path.is_file()
    assert by_id["corrupt.jpg"].status is SheetStatus.ERROR and by_id["corrupt.jpg"].message
    assert result.columns == ["q1", "q2", "q3", "q4", "q5"]


def test_progress_and_cancellation(workspace, library, samples, tmp_path):
    for i in range(3):
        (tmp_path / f"s{i}.png").write_bytes((samples / "answer-key" / "using-csv" / "adrian_omr.png").read_bytes())
    sheets = collect_sheets([tmp_path])
    seen = []
    result = _run(workspace, library, samples, sheets, hooks=EngineHooks(on_progress=lambda d, n: seen.append((d, n))))
    assert seen[-1] == (3, 3) and len(result.sheets) == 3

    template = library.list()[0]
    store = SessionStore(workspace.sessions)
    session = store.create(template, sheets, RunOptions())
    with pytest.raises(RunCancelled):
        OMREngine().execute(session, RunOptions(), EngineHooks(is_cancelled=lambda: True))


def test_export_excel_and_csv(workspace, library, samples, tmp_path):
    sheets = collect_sheets([samples / "answer-key" / "using-csv" / "adrian_omr.png"])
    result = _run(workspace, library, samples, sheets)
    xlsx = ResultsExporter().export_excel(result, tmp_path / "out")
    wb = load_workbook(xlsx)
    assert wb.sheetnames == ["Summary", "Results", "Needs Review"]
    header = [c.value for c in wb["Results"][1]]
    assert header[:3] == ["File", "Status", "Score"]
    assert wb["Results"]["C2"].value == 5.0
    csv_path = ResultsExporter().export_csv(result, tmp_path / "out")
    assert csv_path.read_text(encoding="utf-8-sig").splitlines()[0].startswith("File,Status,Score")


def test_excel_cells_never_become_formulas(workspace, library, samples, tmp_path):
    result = _run(workspace, library, samples, collect_sheets([samples / "answer-key" / "using-csv" / "adrian_omr.png"]))
    result.sheets[0].file_id = "=HYPERLINK(\"http://evil\")"
    wb = load_workbook(ResultsExporter().export_excel(result, tmp_path / "safe"))
    assert wb["Results"]["A2"].data_type == "s"
