"""Excel / CSV export of a processing run."""
from __future__ import annotations

import csv
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..models import RunResult, SheetResult, SheetStatus

_HEADER_FILL = PatternFill("solid", fgColor="0F172A")
_HEADER_FONT = Font(bold=True, color="FFFFFF", name="Calibri")
_STATUS_FILL = {
    SheetStatus.PROCESSED: PatternFill("solid", fgColor="DCFCE7"),
    SheetStatus.MULTI_MARKED: PatternFill("solid", fgColor="FEF3C7"),
    SheetStatus.ERROR: PatternFill("solid", fgColor="FEE2E2"),
}
_THIN = Side(style="thin", color="E2E8F0")
_BORDER = Border(left=_THIN, right=_THIN, top=_THIN, bottom=_THIN)


def _text(ws, row: int, col: int, value):
    """Write a value; force text for strings so cells can never become formulas."""
    cell = ws.cell(row=row, column=col, value=value)
    if isinstance(value, str):
        cell.data_type = "s"
    cell.border = _BORDER
    return cell


def _header(ws, headers: list[str]) -> None:
    for col, title in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=col, value=title)
        cell.fill, cell.font, cell.border = _HEADER_FILL, _HEADER_FONT, _BORDER
        cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.freeze_panes = "B2"
    ws.row_dimensions[1].height = 22


def _autofit(ws, minimum: int = 8, maximum: int = 48) -> None:
    for column in ws.columns:
        width = max((len(str(c.value)) for c in column if c.value is not None), default=0)
        ws.column_dimensions[get_column_letter(column[0].column)].width = min(max(width + 3, minimum), maximum)


def _rows(result: RunResult, sheets: list[SheetResult], with_message: bool):
    headers = ["File", "Status"] + (["Score"] if result.session.graded else []) + result.columns
    if with_message:
        headers.append("Reason")
    for s in sheets:
        row: list = [s.file_id, s.status.value]
        if result.session.graded:
            row.append(s.score)
        row += [s.answers.get(c, "") for c in result.columns]
        if with_message:
            row.append(s.message)
        yield headers, row, s


class ResultsExporter:
    def export_excel(self, result: RunResult, destination: Path) -> Path:
        wb = Workbook()
        self._summary_sheet(wb.active, result)
        self._data_sheet(wb.create_sheet("Results"), result, result.sheets, with_message=False)
        review = [s for s in result.sheets if s.status is not SheetStatus.PROCESSED]
        self._data_sheet(wb.create_sheet("Needs Review"), result, review, with_message=True)
        destination = Path(destination).with_suffix(".xlsx")
        destination.parent.mkdir(parents=True, exist_ok=True)
        wb.save(destination)
        return destination

    def export_csv(self, result: RunResult, destination: Path) -> Path:
        destination = Path(destination).with_suffix(".csv")
        destination.parent.mkdir(parents=True, exist_ok=True)
        with destination.open("w", newline="", encoding="utf-8-sig") as handle:  # BOM keeps Excel happy
            writer = csv.writer(handle)
            wrote_header = False
            for headers, row, _ in _rows(result, result.sheets, with_message=False):
                if not wrote_header:
                    writer.writerow(headers)
                    wrote_header = True
                writer.writerow(row)
            if not wrote_header:
                writer.writerow(["File", "Status"] + result.columns)
        return destination

    # ---- sheets ---------------------------------------------------------- #
    @staticmethod
    def _data_sheet(ws, result: RunResult, sheets: list[SheetResult], with_message: bool) -> None:
        wrote_header = False
        for r, (headers, row, sheet) in enumerate(_rows(result, sheets, with_message), start=2):
            if not wrote_header:
                _header(ws, headers)
                wrote_header = True
            for c, value in enumerate(row, start=1):
                cell = _text(ws, r, c, value)
                if c == 2:
                    cell.fill = _STATUS_FILL[sheet.status]
                if headers[c - 1] == "Score" and isinstance(value, float):
                    cell.number_format = "0.00"
        if not wrote_header:
            _header(ws, ["File", "Status"] + result.columns + (["Reason"] if with_message else []))
        ws.auto_filter.ref = ws.dimensions
        _autofit(ws)

    @staticmethod
    def _summary_sheet(ws, result: RunResult) -> None:
        ws.title = "Summary"
        info = result.session
        ws["A1"] = "OMR Results Summary"
        ws["A1"].font = Font(bold=True, size=16, color="0F172A")
        facts: list[tuple[str, object]] = [
            ("Template", info.template_name),
            ("Processed on", info.created_at.strftime("%d %b %Y %H:%M")),
            ("Exported on", datetime.now().strftime("%d %b %Y %H:%M")),
            ("Total sheets", len(result.sheets)),
            ("Processed", result.count(SheetStatus.PROCESSED)),
            ("Multi-marked (review)", result.count(SheetStatus.MULTI_MARKED)),
            ("Errors", result.count(SheetStatus.ERROR)),
        ]
        if info.graded and result.scores:
            facts += [("Average score", round(result.average_score, 2)),
                      ("Highest score", max(result.scores)), ("Lowest score", min(result.scores))]
        for r, (label, value) in enumerate(facts, start=3):
            ws.cell(row=r, column=1, value=label).font = Font(bold=True)
            _text(ws, r, 2, value)
        ws.column_dimensions["A"].width = 26
        ws.column_dimensions["B"].width = 34
